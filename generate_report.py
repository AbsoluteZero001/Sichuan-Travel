from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)

title = doc.add_heading('四川旅游景点推荐系统', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('项目实践报告')
doc.add_paragraph('')

heading1 = doc.add_heading('一、项目概述', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

doc.add_paragraph('本项目是一个基于Flask框架的四川旅游景点推荐系统，采用前后端分离架构，提供景点浏览、搜索、收藏、AI问答等功能。')
doc.add_paragraph('')

heading1 = doc.add_heading('二、技术栈', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

table = doc.add_table(rows=3, cols=2)
table.cell(0, 0).text = '前端'
table.cell(0, 1).text = 'HTML5 + CSS3 + JavaScript'
table.cell(1, 0).text = '后端'
table.cell(1, 1).text = 'Python Flask'
table.cell(2, 0).text = '数据库'
table.cell(2, 1).text = 'MySQL'
doc.add_paragraph('')

heading1 = doc.add_heading('三、数据库设计', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

doc.add_heading('3.1 数据库连接配置', level=2)
doc.add_paragraph('文件路径：backend/sjk.py')
code = '''import mysql.connector
from mysql.connector import Error

config = {
    'host': 'localhost',
    'user': 'root',
    'password': 'Root@Abc456',
    'database': 'sichuan_tourism',
    'charset': 'utf8mb4'
}

def get_connection():
    try:
        conn = mysql.connector.connect(**config)
        return conn
    except Error as e:
        print(f"数据库连接失败: {e}")
        return None'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 导入mysql.connector库用于MySQL数据库连接')
doc.add_paragraph('2. config字典配置数据库连接参数：主机、用户名、密码、数据库名、字符集')
doc.add_paragraph('3. get_connection()函数尝试建立数据库连接，失败时返回None')
doc.add_paragraph('')

doc.add_heading('3.2 数据表结构', level=2)
doc.add_paragraph('文件路径：backend/sjk.py')
code = '''def init_db():
    conn = get_connection()
    if conn:
        cursor = conn.cursor()
        
        # 创建用户表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT PRIMARY KEY AUTO_INCREMENT,
                username VARCHAR(50) UNIQUE NOT NULL,
                password VARCHAR(255) NOT NULL,
                role ENUM('user', 'admin') DEFAULT 'user',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """)
        
        # 创建景点表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS spots (
                id INT PRIMARY KEY AUTO_INCREMENT,
                name VARCHAR(100) NOT NULL,
                city VARCHAR(50) NOT NULL,
                season ENUM('春季','夏季','秋季','冬季'),
                price DECIMAL(10,2) DEFAULT 0,
                open_time VARCHAR(50),
                image VARCHAR(255),
                images TEXT,
                description TEXT,
                guide TEXT,
                tips TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """)
        
        # 创建收藏表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS spot_favorites (
                id INT PRIMARY KEY AUTO_INCREMENT,
                spot_id INT NOT NULL,
                user_id INT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (spot_id) REFERENCES spots(id) ON DELETE CASCADE,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
                UNIQUE KEY unique_favorite (spot_id, user_id)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """)
        
        conn.commit()
        conn.close()'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. users表：存储用户信息，包含用户名、密码、角色字段')
doc.add_paragraph('2. spots表：存储景点信息，包含名称、城市、季节、价格、图片、描述等字段')
doc.add_paragraph('3. spot_favorites表：存储用户收藏关系，通过外键关联spots和users表')
doc.add_paragraph('4. UNIQUE KEY unique_favorite约束确保用户对同一景点只能收藏一次')
doc.add_paragraph('')

heading1 = doc.add_heading('四、后端代码', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

doc.add_heading('4.1 主应用入口', level=2)
doc.add_paragraph('文件路径：backend/app.py')
code = '''from flask import Flask, send_from_directory, request, jsonify
from flask_cors import CORS
from routes.users import users_bp, verify_token
from routes.spots import spots_bp
from routes.favorites import favorites_bp
from routes.ai import ai_bp
from sjk import init_db, get_connection
import os
import json

app = Flask(__name__)
CORS(app, origins='*')

# 注册蓝图
app.register_blueprint(users_bp, url_prefix='/api/users')
app.register_blueprint(spots_bp, url_prefix='/api/spots')
app.register_blueprint(favorites_bp, url_prefix='/api/favorites')
app.register_blueprint(ai_bp, url_prefix='/api/ai')

if __name__ == '__main__':
    init_db()
    app.run(host='0.0.0.0', port=3000, debug=False)'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 创建Flask应用实例，配置CORS跨域支持')
doc.add_paragraph('2. register_blueprint注册各功能模块的路由蓝图')
doc.add_paragraph('3. 启动时调用init_db()初始化数据库')
doc.add_paragraph('4. 监听0.0.0.0:3000端口启动服务')
doc.add_paragraph('')

doc.add_heading('4.2 景点模糊搜索接口', level=2)
doc.add_paragraph('文件路径：backend/app.py')
code = '''@app.route('/api/spots/', methods=['GET'])
def get_spots():
    keyword = request.args.get('keyword', '')
    season = request.args.get('season', '')
    
    conn = get_connection()
    if not conn:
        return jsonify({'success': False, 'message': '数据库连接失败'}), 500
    
    try:
        cursor = conn.cursor(dictionary=True)
        
        query = """
            SELECT s.*, COALESCE(f.favorite_count, 0) as favorite_count
            FROM spots s
            LEFT JOIN (
                SELECT spot_id, COUNT(*) as favorite_count
                FROM spot_favorites
                GROUP BY spot_id
            ) f ON s.id = f.spot_id
            WHERE 1=1
        """
        params = []
        
        if keyword:
            query += " AND (s.name LIKE %s OR s.city LIKE %s OR s.season LIKE %s)"
            params.append(f'%{keyword}%')
            params.append(f'%{keyword}%')
            params.append(f'%{keyword}%')
        
        if season:
            query += " AND s.season = %s"
            params.append(season)
        
        query += " ORDER BY f.favorite_count DESC, s.id ASC"
        
        cursor.execute(query, params)
        spots = cursor.fetchall()
        
        return jsonify({'success': True, 'data': spots})
    finally:
        conn.close()'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 获取URL参数keyword和season')
doc.add_paragraph('2. 使用LEFT JOIN统计每个景点的收藏数量（COALESCE处理无收藏情况）')
doc.add_paragraph('3. keyword参数同时匹配name、city、season三个字段，实现模糊搜索')
doc.add_paragraph('4. 结果按收藏数量降序排列')
doc.add_paragraph('')

doc.add_heading('4.3 用户认证模块', level=2)
doc.add_paragraph('文件路径：backend/routes/users.py')
code = '''import jwt
from datetime import datetime, timedelta

SECRET_KEY = 'sichuan_tourism_secret_key'

def generate_token(user_id, username, role):
    payload = {
        'user_id': user_id,
        'username': username,
        'role': role,
        'exp': datetime.utcnow() + timedelta(hours=24)
    }
    return jwt.encode(payload, SECRET_KEY, algorithm='HS256')

def verify_token(token):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=['HS256'])
        return payload
    except jwt.ExpiredSignatureError:
        return None
    except jwt.InvalidTokenError:
        return None'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 使用JWT（JSON Web Token）实现无状态认证')
doc.add_paragraph('2. generate_token()生成包含用户信息的token，有效期24小时')
doc.add_paragraph('3. verify_token()验证token有效性，过期或无效时返回None')
doc.add_paragraph('')

code = '''@users_bp.route('/login', methods=['POST'])
def login():
    data = request.get_json()
    username = data.get('username')
    password = data.get('password')
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE username = %s AND password = %s", (username, password))
    user = cursor.fetchone()
    
    if not user:
        return jsonify({'success': False, 'message': '用户名或密码错误'}), 401
    
    token = generate_token(user[0], user[1], user[3])
    
    return jsonify({
        'success': True,
        'token': token,
        'user': {'id': user[0], 'username': user[1], 'role': user[3]}
    })'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 接收POST请求的用户名和密码')
doc.add_paragraph('2. 查询数据库验证用户身份')
doc.add_paragraph('3. 验证成功后生成JWT token返回给前端')
doc.add_paragraph('')

doc.add_heading('4.4 收藏功能模块', level=2)
doc.add_paragraph('文件路径：backend/routes/favorites.py')
code = '''@favorites_bp.route('/', methods=['POST'])
def add_favorite():
    auth_header = request.headers.get('Authorization')
    if not auth_header or not auth_header.startswith('Bearer '):
        return jsonify({'success': False, 'message': '未登录'}), 401
    
    token = auth_header.split(' ')[1]
    payload = verify_token(token)
    if not payload:
        return jsonify({'success': False, 'message': '登录已过期'}), 401
    
    data = request.get_json()
    spot_id = data.get('spot_id')
    
    conn = get_connection()
    cursor = conn.cursor()
    
    try:
        cursor.execute("INSERT INTO spot_favorites (spot_id, user_id) VALUES (%s, %s)", 
                       (spot_id, payload['user_id']))
        conn.commit()
        return jsonify({'success': True, 'message': '收藏成功'})
    except Exception as e:
        return jsonify({'success': False, 'message': '已收藏'}), 400'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 从请求头获取Authorization Bearer token')
doc.add_paragraph('2. 验证token有效性')
doc.add_paragraph('3. 将收藏记录插入spot_favorites表')
doc.add_paragraph('4. 由于数据库有UNIQUE约束，重复收藏会抛出异常')
doc.add_paragraph('')

doc.add_heading('4.5 AI问答模块', level=2)
doc.add_paragraph('文件路径：backend/routes/ai.py')
code = '''@ai_bp.route('/chat', methods=['POST'])
def chat():
    data = request.get_json()
    question = data.get('question', '')
    
    qa_pairs = [
        {
            'keywords': ['季节', '什么时候去', '几月'],
            'answer': '四川气候多样，不同季节有不同的美景...'
        },
        {
            'keywords': ['交通', '怎么去', '机场', '高铁'],
            'answer': '四川省交通便利，主要交通方式包括...'
        },
        {
            'keywords': ['美食', '火锅', '川菜'],
            'answer': '四川美食以麻辣鲜香著称...'
        }
    ]
    
    for qa in qa_pairs:
        if any(keyword in question for keyword in qa['keywords']):
            return jsonify({'success': True, 'answer': qa['answer']})
    
    return jsonify({'success': True, 'answer': '抱歉，我暂时无法回答这个问题。'})'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 接收用户提问')
doc.add_paragraph('2. 遍历预定义的问答对，匹配关键词')
doc.add_paragraph('3. 返回匹配到的答案，支持季节、交通、美食等多方面问答')
doc.add_paragraph('')

heading1 = doc.add_heading('五、前端代码', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

doc.add_heading('5.1 景点列表加载', level=2)
doc.add_paragraph('文件路径：frontend/index.html')
code = '''let currentKeyword = '';

async function loadSpots(keyword = '') {
    currentKeyword = keyword;
    const response = await fetch(`/api/spots/?keyword=${encodeURIComponent(keyword)}`);
    const data = await response.json();
    
    if (data.success) {
        renderSpots(data.data);
    }
}

async function renderSpots(spots) {
    const container = document.getElementById('spots-container');
    container.innerHTML = '';
    
    for (const spot of spots) {
        const isFavorite = await checkFavorite(spot.id);
        
        const card = document.createElement('div');
        card.className = 'spot-card';
        card.innerHTML = `
            <img src="${spot.image}" alt="${spot.name}">
            <div class="card-content">
                <h3>${spot.name}</h3>
                <p>${spot.city} · ${spot.season}</p>
                <button onclick="handleCardFavorite(${spot.id}, this)" 
                        class="btn ${isFavorite ? 'filled' : ''}">
                    ${isFavorite ? '♥ 已收藏' : '♡ 收藏'}
                </button>
            </div>
        `;
        container.appendChild(card);
    }
}'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. loadSpots()通过fetch调用后端API获取景点数据')
doc.add_paragraph('2. renderSpots()动态生成景点卡片HTML')
doc.add_paragraph('3. 根据收藏状态显示不同的爱心样式')
doc.add_paragraph('')

doc.add_heading('5.2 收藏功能实现', level=2)
doc.add_paragraph('文件路径：frontend/index.html')
code = '''async function handleCardFavorite(spotId, btn) {
    const token = localStorage.getItem('token');
    if (!token) {
        alert('请先登录');
        window.location.href = 'login.html';
        return;
    }
    
    const response = await fetch('/api/favorites/', {
        method: 'POST',
        headers: {
            'Content-Type': 'application/json',
            'Authorization': `Bearer ${token}`
        },
        body: JSON.stringify({ spot_id: spotId })
    });
    
    const data = await response.json();
    
    if (data.success) {
        btn.innerHTML = '♥ 已收藏';
        btn.classList.add('filled');
        loadSpots(currentKeyword);
    } else {
        alert(data.message);
    }
}'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)
doc.add_paragraph('')
doc.add_paragraph('【代码解释】')
doc.add_paragraph('1. 从localStorage获取登录token')
doc.add_paragraph('2. 携带Authorization请求头发送收藏请求')
doc.add_paragraph('3. 成功后更新按钮状态并重新加载景点列表')
doc.add_paragraph('')

heading1 = doc.add_heading('六、API接口列表', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

table = doc.add_table(rows=8, cols=3)
table.cell(0, 0).text = '接口路径'
table.cell(0, 1).text = '方法'
table.cell(0, 2).text = '功能描述'
table.cell(1, 0).text = '/api/users/register'
table.cell(1, 1).text = 'POST'
table.cell(1, 2).text = '用户注册'
table.cell(2, 0).text = '/api/users/login'
table.cell(2, 1).text = 'POST'
table.cell(2, 2).text = '用户登录'
table.cell(3, 0).text = '/api/spots/'
table.cell(3, 1).text = 'GET'
table.cell(3, 2).text = '获取景点列表（支持模糊搜索）'
table.cell(4, 0).text = '/api/spots/<id>'
table.cell(4, 1).text = 'GET'
table.cell(4, 2).text = '获取单个景点详情'
table.cell(5, 0).text = '/api/favorites/'
table.cell(5, 1).text = 'POST'
table.cell(5, 2).text = '添加收藏'
table.cell(6, 0).text = '/api/favorites/<spot_id>'
table.cell(6, 1).text = 'DELETE'
table.cell(6, 2).text = '取消收藏'
table.cell(7, 0).text = '/api/ai/chat'
table.cell(7, 1).text = 'POST'
table.cell(7, 2).text = 'AI问答'

doc.add_paragraph('')

heading1 = doc.add_heading('七、项目运行', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

doc.add_paragraph('1. 安装依赖：')
code = 'pip install flask flask-cors mysql-connector-python pyjwt'
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)

doc.add_paragraph('')
doc.add_paragraph('2. 启动服务：')
code = 'cd backend && python app.py'
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)

doc.add_paragraph('')
doc.add_paragraph('3. 访问地址：http://localhost:3000')

doc.add_paragraph('')
heading1 = doc.add_heading('八、项目结构', level=1)
heading1.style.font.color.rgb = RGBColor(0, 0, 128)

code = '''Sichuan Travel/
├── backend/
│   ├── app.py              # 主应用入口
│   ├── sjk.py              # 数据库配置与初始化
│   └── routes/
│       ├── users.py        # 用户模块路由
│       ├── spots.py        # 景点模块路由
│       ├── favorites.py    # 收藏模块路由
│       └── ai.py           # AI问答模块路由
├── frontend/
│   ├── index.html          # 首页
│   ├── about-sichuan.html  # 关于四川
│   ├── all-spots.html      # 全部景点
│   ├── spot-detail.html    # 景点详情
│   ├── travel-tips.html    # 旅游指南
│   ├── login.html          # 登录页
│   ├── register.html       # 注册页
│   ├── admin-panel.html    # 管理员面板
│   ├── user-profile.html   # 用户中心
│   ├── css/
│   │   └── style.css       # 样式文件
│   ├── js/
│   │   └── common.js       # 公共脚本
│   └── images/             # 图片资源
└── README.md'''
p = doc.add_paragraph()
run = p.add_run(code)
run.font.name = 'Consolas'
run.font.size = Pt(9)

doc.save('四川旅游景点推荐系统-项目实践报告.docx')
print('文档已保存')
